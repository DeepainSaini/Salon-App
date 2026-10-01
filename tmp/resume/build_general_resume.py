from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_resume import (
    BLACK,
    MUTED,
    ROOT,
    add_hyperlink,
    add_project,
    add_section_heading,
    configure_styles,
    set_cell_free_font,
)


OUTPUT = ROOT / "output" / "docx" / "Deepain_Saini_Full_Stack_Developer_Resume.docx"


def build_resume():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_before = Pt(0)
    name.paragraph_format.space_after = Pt(3.5)
    name.paragraph_format.keep_with_next = True
    set_cell_free_font(name.add_run("DEEPAIN SAINI"), 21.5, bold=True, color=BLACK)

    contact = doc.add_paragraph(style="Contact Line")
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_hyperlink(contact, "deepainsaini111@gmail.com", "mailto:deepainsaini111@gmail.com", size=9.7)
    set_cell_free_font(contact.add_run("  |  +91 96499 90222  |  "), 9.7, color=MUTED)
    add_hyperlink(contact, "github.com/DeepainSaini", "https://github.com/DeepainSaini", size=9.7)

    add_section_heading(doc, "Professional Summary")
    summary = doc.add_paragraph()
    summary.paragraph_format.space_after = Pt(6)
    summary.paragraph_format.line_spacing = 1.08
    set_cell_free_font(
        summary.add_run(
            "Full-stack developer with hands-on experience building and deploying backend-focused web applications using JavaScript, Node.js, Express.js, MySQL, Sequelize, and AWS. Built REST APIs, authentication, database transactions, real-time messaging, payment workflows, file storage, background jobs, logging, and API validation. Comfortable using Git/GitHub branch, pull, push, and merge workflows, breaking features into tasks, testing APIs with Postman, and learning quickly across ambiguous requirements."
        ),
        9.8,
    )

    add_section_heading(doc, "Education and Training")
    education = doc.add_paragraph()
    education.paragraph_format.space_after = Pt(2.5)
    education.paragraph_format.keep_with_next = True
    set_cell_free_font(education.add_run("Bachelor of Technology in Computer Science and Engineering"), 10.05, bold=True)
    set_cell_free_font(education.add_run("  |  Graduated: 2021"), 9.8, bold=True, color=MUTED)
    university = doc.add_paragraph()
    university.paragraph_format.space_after = Pt(3.5)
    set_cell_free_font(university.add_run("Guru Jambheshwar University of Science and Technology, Hisar"), 9.65, color=MUTED)

    training = doc.add_paragraph()
    training.paragraph_format.space_after = Pt(4.5)
    set_cell_free_font(training.add_run("Full Stack Development Program - Sharpener Tech"), 10.0, bold=True)
    set_cell_free_font(
        training.add_run("  |  1-year project-based training in JavaScript, Node.js, Express.js, MySQL, AWS, Git/GitHub"),
        9.55,
        color=MUTED,
    )

    add_section_heading(doc, "Technical Skills")
    skill_lines = [
        ("Programming", "JavaScript (ES6+), SQL, HTML, CSS"),
        ("Backend", "Node.js, Express.js, REST APIs, Socket.IO, JWT, bcrypt, scheduled jobs"),
        ("Databases", "MySQL, Sequelize ORM, migrations, transactions, pagination, indexing"),
        ("Testing and Quality", "Postman API testing, input validation, error handling, Winston logging"),
        ("Cloud and DevOps", "AWS S3, AWS RDS, AWS EC2, Nginx, PM2, Jenkins CI/CD"),
        ("Tools and Collaboration", "Git, GitHub, branching, pull/push/merge workflows, task management, Google Workspace"),
    ]
    for label, value in skill_lines:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3.5)
        set_cell_free_font(paragraph.add_run(f"{label}: "), 9.6, bold=True)
        set_cell_free_font(paragraph.add_run(value), 9.6)

    add_section_heading(doc, "Software Development Projects")

    add_project(
        doc,
        "Salon Appointment Platform",
        "https://github.com/DeepainSaini/Salon-App",
        "Jul 2026 - Present",
        "JavaScript, Node.js, Express.js, MySQL, Sequelize, JWT, Cashfree, Brevo",
        [
            "Built customer, staff, and admin workflows for appointment booking, dynamic slot calculation, availability and conflict checks, rescheduling, cancellation, reviews, and role-based access.",
            "Integrated JWT/cookie authentication, Cashfree sandbox payments, Brevo confirmation and reminder emails, and Sequelize transactions for multi-step booking and payment flows.",
        ],
    )

    add_project(
        doc,
        "Group Chat Application",
        "https://github.com/DeepainSaini/group-chat-app",
        "Sep 2025 - May 2026",
        "JavaScript, Node.js, Express.js, Socket.IO, MySQL, Sequelize, AWS S3",
        [
            "Developed authenticated personal and group messaging with Socket.IO rooms, persistent MySQL storage, and server-side membership checks before group-message delivery.",
            "Added AWS S3 media sharing and a scheduled archival job that moves messages older than 24 hours to a separate table to keep active-message queries focused.",
        ],
    )

    add_project(
        doc,
        "Daily Expense Tracker",
        "https://github.com/DeepainSaini/daily-expense-tracking",
        "May 2025 - Sep 2025",
        "JavaScript, Node.js, Express.js, MySQL, Sequelize, AWS, Cashfree, JWT",
        [
            "Implemented expense tracking with JWT/bcrypt authentication, pagination, transactional create/delete flows, password-reset emails, premium payments, and AWS S3 report exports.",
            "Deployed on AWS EC2 with MySQL on RDS, PM2, Jenkins CI/CD, and Winston logging; optimized leaderboard reads by maintaining per-user expense totals.",
        ],
    )

    doc.core_properties.title = "Deepain Saini - Full Stack Developer Resume"
    doc.core_properties.subject = "General full-stack and backend developer resume"
    doc.core_properties.author = "Deepain Saini"
    doc.core_properties.keywords = (
        "full stack developer, backend developer, JavaScript, Node.js, Express.js, MySQL, Sequelize, REST APIs, "
        "Socket.IO, AWS, Nginx, PM2, Jenkins, Git, Postman, CI/CD"
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_resume()
